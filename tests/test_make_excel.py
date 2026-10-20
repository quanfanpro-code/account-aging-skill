# -*- coding: utf-8 -*-
"""make_excel 输出层测试"""
import pandas as pd
from openpyxl import load_workbook
from make_excel import _infer_column_type, make_excel


def test_空DataFrame也建Sheet并写占位行(tmp_path):
    """空 DataFrame 不应被静默跳过：建 sheet 并写"（无数据）"占位"""
    out = tmp_path / "empty.xlsx"
    make_excel(
        [("有数据", pd.DataFrame({"金额": [1.0]})),
         ("空表", pd.DataFrame({"金额": []}))],
        str(out),
    )
    wb = load_workbook(out)
    assert "有数据" in wb.sheetnames
    assert "空表" in wb.sheetnames
    assert wb["有数据"]["B1"].alignment.horizontal == "center"
    texts = [str(c.value) for row in wb["空表"].iter_rows() for c in row if c.value]
    assert any("无数据" in t for t in texts)


def test_列类型推断_小额金额不误判百分比():
    """0~1 之间的浮点金额（如万元单位小额数据）仍应识别为金额而非百分比"""
    series = pd.Series([0.12, 0.35, 0.78, 1.0])
    assert _infer_column_type("借方金额", series) == "money"
    assert _infer_column_type("期初余额", series) == "money"


def test_列类型推断_名称关键词优先():
    assert _infer_column_type("毛利率", pd.Series([0.12, 0.35])) == "pct"
    assert _infer_column_type("业务日期", pd.Series(["2025-01-01"])) == "date"
    assert _infer_column_type("科目编码", pd.Series(["1122", "2202"])) == "text"
    assert _infer_column_type("账龄天数", pd.Series([30, 90, 180])) == "number"
