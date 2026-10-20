"""截止日账龄算完后，同一单位成对科目的展示对抵。"""
from datetime import date

from openpyxl import load_workbook
import pytest

from src.core.models import AgingResult
from src.core.report import ReportGenerator


def result(code, name, unit, direction, balance, buckets):
    return AgingResult(
        subject_code=code, subject_name=name, detail_name=unit, direction=direction,
        opening_balance=balance, total_debit=0.0, total_credit=0.0,
        closing_balance=balance, bucket_totals=buckets,
        negative_balance=min(balance, 0.0),
    )


CONFIG = {"cutoff_date": date(2025, 12, 31), "buckets": [(0, 365), (365, 9999)]}


@pytest.mark.parametrize("asset_code,asset_name,liability_code,liability_name", [
    ("1122", "应收账款", "2203", "预收账款"),
    ("1123", "预付账款", "2202", "应付账款"),
    ("1221", "其他应收款", "2241", "其他应付款"),
    ("1486", "合同资产", "2204", "合同负债"),
])
def test_四对科目同一单位在输出阶段对抵_先冲老账(asset_code, asset_name, liability_code, liability_name):
    asset = result(asset_code, asset_name, "单位A", "debit", 100.0, [40.0, 60.0])
    liability = result(liability_code, liability_name, "单位A", "credit", 30.0, [30.0, 0.0])
    frame = ReportGenerator()._build_paired_offset([asset, liability], CONFIG)
    row = frame.iloc[0]
    assert len(frame) == 1
    assert row["对抵金额"] == pytest.approx(30.0)
    assert row["对抵后方向"] == "借方"
    assert row["对抵后余额"] == pytest.approx(70.0)
    assert row["0-365天"] == pytest.approx(40.0)
    assert row["365天以上"] == pytest.approx(30.0)
    assert row["分档勾稽差异"] == pytest.approx(0.0)
    assert asset.bucket_totals == [40.0, 60.0]  # 原账龄结果不改
    assert liability.bucket_totals == [30.0, 0.0]


def test_同一单位不同组合不相互对抵():
    receivable = result("1122", "应收账款", "单位A", "debit", 100.0, [40.0, 60.0])
    payable = result("2202", "应付账款", "单位A", "credit", 30.0, [30.0, 0.0])
    frame = ReportGenerator()._build_paired_offset([receivable, payable], CONFIG)
    assert len(frame) == 2
    assert frame["对抵金额"].tolist() == [0.0, 0.0]


def test_负债侧余额较大时保留贷方账龄():
    asset = result("1122", "应收账款", "单位A", "debit", 30.0, [30.0, 0.0])
    liability = result("2203", "预收账款", "单位A", "credit", 100.0, [40.0, 60.0])
    row = ReportGenerator()._build_paired_offset([asset, liability], CONFIG).iloc[0]
    assert row["对抵金额"] == pytest.approx(30.0)
    assert row["对抵后方向"] == "贷方"
    assert row["对抵后余额"] == pytest.approx(70.0)
    assert row["0-365天"] == pytest.approx(40.0)
    assert row["365天以上"] == pytest.approx(30.0)


def test_空明细及反向余额不自动对抵():
    items = [
        result("1122", "应收账款", "（未填明细）", "debit", 100.0, [100.0, 0.0]),
        result("2203", "预收账款", "（未填明细）", "credit", 30.0, [30.0, 0.0]),
        result("1123", "预付账款", "供应商B", "debit", -20.0, [0.0, 0.0]),
        result("2202", "应付账款", "供应商B", "credit", 30.0, [30.0, 0.0]),
    ]
    frame = ReportGenerator()._build_paired_offset(items, CONFIG)
    assert frame["对抵金额"].tolist() == [0.0, 0.0]
    assert all("未自动对抵" in status for status in frame["处理状态"])
    assert frame["对抵后余额"].isna().all()


def test_对抵直接进入唯一最终结果_依据单独留痕(tmp_path):
    from dataclasses import replace
    asset=result('1122','应收账款','单位A','debit',100,[40,60])
    liability=result('2203','预收账款','单位A','credit',30,[30,0])
    final=replace(asset,closing_balance=70,bucket_totals=[40,30])
    target=tmp_path/'最终账龄.xlsx'
    ReportGenerator().generate(target,[final],{**CONFIG,'source_results':[asset,liability],'offset_records':[{'往来组合':'应收账款 / 预收账款','单位':'单位A','资产侧余额':100,'负债侧余额':30,'对抵金额':30,'最终借方净额':70}]})
    wb=load_workbook(target,data_only=True)
    assert wb['账龄计算结果']['G2'].value==70
    assert wb['账龄计算结果']['E2'].value==40 and wb['账龄计算结果']['F2'].value==30
    assert wb['对抵记录']['F2'].value==30
    assert '成对对抵后账龄' not in wb.sheetnames
    wb.close()
