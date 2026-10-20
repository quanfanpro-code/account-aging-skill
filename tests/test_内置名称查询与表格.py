import hashlib
import importlib.util
from pathlib import Path
import subprocess
import sys
import json

import openpyxl
import pytest

from src.company_registry.client import CompanyMatch,CompanyDetail,Shareholder,CicpaError,CicpaClient
from src.name_matching.整理表格 import 读取名称列,写出结果
from src.name_matching.归一引擎 import 归一分档


def query_module():
    spec=importlib.util.spec_from_file_location('registry_query',Path(__file__).resolve().parents[1]/'scripts/查询名称登记.py')
    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module);return module


def test_登记查询保留真实身份及具体来源():
    class Client:
        def search_companies(self,name,page_size):return [CompanyMatch('公司甲','成都甲有限公司')]
        def get_company_detail(self,identifier):return CompanyDetail(identifier,'成都甲有限公司','91330000AAAAAAAAAA')
        def get_shareholders(self,identifier):return [Shareholder('股东甲',60)]
    result=query_module().lookup_name('成都甲',Client())
    assert result['matches'][0]['credit_code']=='91330000AAAAAAAAAA'
    assert 'find_company_basic_info?orgid=' in result['matches'][0]['source']
    assert result['matches'][0]['queried_at']
    assert result['matches'][0]['shareholders'][0]['name']=='股东甲'


def test_最短路线只查搜索详情股东():
    calls=[]
    class Client:
        def search_companies(self,name,page_size):calls.append('搜索');return [CompanyMatch('甲','成都甲有限公司')]
        def get_company_detail(self,identifier):calls.append('详情');return CompanyDetail(identifier,'成都甲有限公司','91330000AAAAAAAAAA')
        def get_shareholders(self,identifier):calls.append('股东');return []
    result=query_module().lookup_name('成都甲有限公司',Client())
    assert result['status']=='found' and calls==['搜索','详情','股东']


def test_多个搜索主体不先遍查所有详情股东():
    class Client:
        def search_companies(self,name,page_size):return [CompanyMatch('甲','成都甲有限公司'),CompanyMatch('乙','成都乙有限公司')]
        def get_company_detail(self,identifier):raise AssertionError('尚未选定主体，不应查询详情')
    assert query_module().lookup_name('成都',Client())['status']=='choose_company'


def test_股东未取得如实记录但保留已取得详情():
    class Client:
        def search_companies(self,name,page_size):return [CompanyMatch('甲','成都甲有限公司')]
        def get_company_detail(self,identifier):return CompanyDetail(identifier,'成都甲有限公司','91330000AAAAAAAAAA')
        def get_shareholders(self,identifier):raise CicpaError('本次未取得权限')
    result=query_module().lookup_name('成都甲有限公司',Client())
    assert result['status']=='partial' and result['matches'][0]['detail_status']=='取得'
    assert result['matches'][0]['shareholders_status']=='未取得' and result['issues']


def test_注协详情和股东的下划线字段可完整读取():
    calls=[]
    class Response:
        status_code=200;headers={}
        def __init__(self,data):self.data=data
        def json(self):return {'status_code':0,'data':self.data}
    class Session:
        def request(self,method,url,**kwargs):
            calls.append(url)
            if 'home_search' in url:return Response({'list':[{'org_id':'甲','name':'成都甲有限公司'}]})
            if 'find_company_basic_info' in url:return Response({'name':'成都甲有限公司',
                'unified_social_credit_code':'91330000AAAAAAAAAA','legal_representative':'张某',
                'operating_status':'存续','corp_address':'成都市'})
            if 'stock_holder_newest' in url:return Response({'list':[{'holder_name':'股东甲','held_ratio':'60.0'}]})
            raise AssertionError('不应请求其他维度或批量接口')
    class Limiter:
        def wait(self,kind):return 0
    result=query_module().lookup_name('成都甲有限公司',CicpaClient(session=Session(),limiter=Limiter()))
    company=result['matches'][0]
    assert result['status']=='found' and len(calls)==3
    assert company['credit_code']=='91330000AAAAAAAAAA' and company['legal_person']=='张某'
    assert company['status']=='存续' and company['address']=='成都市'
    assert company['shareholders'][0]['name']=='股东甲' and company['shareholders'][0]['ratio']==60


def test_没查询到不当作主体不存在或允许合并():
    class Client:
        def search_companies(self,name,page_size):return []
    result=query_module().lookup_name('不确定名称',Client())
    assert result['status']=='no_match' and '不代表主体不存在' in result['说明']


def test_缺浏览器工具如实返回_不依赖外部技能(monkeypatch):
    module=query_module();monkeypatch.setattr(module.shutil,'which',lambda name:None)
    assert module.execute(['成都甲'])['status']=='unavailable'


def test_表格工具处理指定工作表且原文件不变(tmp_path):
    source=tmp_path/'源表.xlsx';book=openpyxl.Workbook();book.active.append(['无关内容'])
    sheet=book.create_sheet('指定表');sheet.append(['供应商','金额']);sheet.append(['YF0001_成都甲有限公司',100])
    book.save(source);book.close();digest=hashlib.sha256(source.read_bytes()).digest()
    names=读取名称列(source,'供应商',表名='指定表');output=tmp_path/'归一结果.xlsx'
    写出结果(source,'供应商',names,归一分档([n for _,n in names]),{},output,表名='指定表')
    result=openpyxl.load_workbook(output,data_only=True)
    assert result['指定表'].cell(2,3).value=='成都甲有限公司'
    assert result['指定表'].cell(2,2).value==100 and result.worksheets[0].cell(1,1).value=='无关内容'
    result.close();assert hashlib.sha256(source.read_bytes()).digest()==digest
    with pytest.raises(FileExistsError):写出结果(source,'供应商',names,归一分档([n for _,n in names]),{},output,表名='指定表')


def test_文本表可读取并另存为工作簿(tmp_path):
    source=tmp_path/'原表.csv';source.write_text('供应商,金额\nYF0001_成都甲有限公司,100\n',encoding='utf-8-sig')
    names=读取名称列(source,'供应商');output=tmp_path/'归一.xlsx'
    写出结果(source,'供应商',names,归一分档([n for _,n in names]),{},output)
    book=openpyxl.load_workbook(output,data_only=True)
    assert book.active.cell(2,3).value=='成都甲有限公司';book.close()


def test_源表和输出路径相同时拒绝(tmp_path):
    source=tmp_path/'表.xlsx'
    with pytest.raises(ValueError,match='不能与原文件相同'):
        写出结果(source,'名称',[],{'候选':[],'自动归一':{}},{},source)
