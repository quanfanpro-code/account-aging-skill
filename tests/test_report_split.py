# -*- coding: utf-8 -*-
"""报告分借/贷 sheet 测试"""
import os
from openpyxl import load_workbook
from src.core.report import ReportGenerator
from src.core.models import AgingResult


def test_资产负债在同一最终结果页列示(tmp_path):
    results=[AgingResult('1122','应收账款','客户A','debit',100,0,0,100,bucket_totals=[100]),AgingResult('2202','应付账款','供应商B','credit',200,0,0,200,bucket_totals=[200])]
    output=tmp_path/'单一结果.xlsx'
    ReportGenerator().generate(output,results,{'buckets':[(0,365)]})
    wb=load_workbook(output,data_only=True);s=wb['账龄计算结果']
    assert s['G2'].value=='负债' and s['G3'].value=='资产'
    assert '借方账龄汇总表' not in wb.sheetnames and '成对对抵后账龄' not in wb.sheetnames
    wb.close()


def test_仅有资产时不再增加空负债账龄页(tmp_path):
    result=AgingResult('1122','应收账款','A','debit',100,0,0,100,bucket_totals=[100])
    output=tmp_path/'仅资产.xlsx'
    ReportGenerator().generate(output,[result],{'buckets':[(0,365)]})
    wb=load_workbook(output,data_only=True)
    assert '贷方账龄汇总表' not in wb.sheetnames
    assert wb['账龄计算结果']['B3'].value=='资产合计'
    wb.close()
