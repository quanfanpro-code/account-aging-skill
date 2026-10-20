from datetime import date
import importlib.util
from pathlib import Path
import hashlib
import json
import shutil
import subprocess
import sys

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.core.名称归一 import normalize_accounts


def inputs(names):
    opening=pd.DataFrame(columns=['科目编码','科目名称','明细名称','借方余额','贷方余额'])
    journal=pd.DataFrame([[date(2023,1,1),'1122','应收账款',n,100,0,'销售',i+2]
        for i,n in enumerate(names)],columns=['日期','科目编码','科目名称','明细名称','借方金额','贷方金额','摘要','原账行号'])
    return opening,journal


def test_格式和简称自动统一且所有非名称字段不变():
    o,j=inputs(['上海语荻　光电有限公司','上海语荻光电有限公司','上海语荻光电科技有限公司'])
    result=normalize_accounts(o,j,{})
    assert result['status']=='ready'
    assert set(result['journal']['明细名称'])=={'上海语荻光电科技有限公司'}
    pd.testing.assert_frame_equal(j.drop(columns=['明细名称']),result['journal'][j.columns].drop(columns=['明细名称']))
    assert len(result['changes'])==2


def test_疑似笔误必须先给具体候选而不是直接并():
    o,j=inputs(['成都博立叶光学仪器有限公司','成都傅立叶光学仪器有限公司'])
    result=normalize_accounts(o,j,{})
    assert result['status']=='needs_name_evidence' and len(result['pending'])==1
    assert result['pending'][0]['涉及金额']==200


def test_有据决定合并后贯穿期初和期间且保留依据():
    a,b='成都博立叶光学仪器有限公司','成都傅立叶光学仪器有限公司'
    o,j=inputs([a,b]);o.loc[0]=['1122','应收账款',a,300,0]
    result=normalize_accounts(o,j,{'normalization_decisions':[{
        'names':[a,b],'canonical':b,'basis':'原始发票与更正记录指向同一结算客户','evidence':'发票核验记录A'}]})
    assert result['status']=='ready' and set(result['opening']['明细名称'])=={b}
    assert set(result['journal']['明细名称'])=={b}
    assert any('发票核验记录A' in r['证据来源'] for r in result['changes'])


def test_登记为两家不同主体时不得套用后缀规则合并():
    a,b='成都示例有限公司','成都示例股份有限公司'
    o,j=inputs([a,b])
    config={'normalization_registry':[
        {'name':a,'registered_name':a,'credit_code':'91330000AAAAAAAAAA','source':'登记详情甲'},
        {'name':b,'registered_name':b,'credit_code':'91330000BBBBBBBBBB','source':'登记详情乙'}]}
    result=normalize_accounts(o,j,config)
    assert result['status']=='ready' and set(result['journal']['明细名称'])=={a,b}


def test_个人与占位名称不做模糊合并_个人后缀可整理():
    o,j=inputs(['张小明','张小朋','张小明（个人）','柏油未列明供应商','光坯未列明供应商'])
    result=normalize_accounts(o,j,{})
    assert result['status']=='ready'
    assert list(result['journal']['明细名称'])==['张小明','张小朋','张小明','柏油未列明供应商','光坯未列明供应商']


def test_物料占位可逐行按真实往来方向处理而非跨方向混用():
    o,j=inputs(['柏油','柏油'])
    j.loc[1,['科目编码','科目名称']]=['2202','应付账款']
    result=normalize_accounts(o,j,{'normalization_row_decisions':[
        {'source':'journal','rows':[2],'canonical':'柏油未列明客户','basis':'摘要仅指物料，未列实际客户'},
        {'source':'journal','rows':[3],'canonical':'柏油未列明供应商','basis':'摘要仅指物料，未列实际供应商'}]})
    assert result['status']=='ready'
    assert list(result['journal']['明细名称'])==['柏油未列明客户','柏油未列明供应商']


def test_候选过多先汇报范围而不落盘():
    o,j=inputs([f'上海示例名称{i:02d}有限公司' for i in range(12)])
    result=normalize_accounts(o,j,{})
    assert result['status']=='needs_name_scope' and result['candidate_count']>50


def test_简称同时符合两个不同主体不得自动随便选一家():
    a,b,c='成都示例光电有限公司','成都示例光电','成都示例光电科技有限公司'
    o,j=inputs([a,b,c])
    result=normalize_accounts(o,j,{'normalization_registry':[
        {'name':a,'registered_name':a,'credit_code':'91330000AAAAAAAAAA','source':'登记甲'},
        {'name':c,'registered_name':c,'credit_code':'91330000BBBBBBBBBB','source':'登记乙'}]})
    assert result['status']=='needs_name_evidence'
    assert all(b in item['名称'] for item in result['pending'])


def test_最终账龄确实使用归一名称并同文件保存对照(tmp_path):
    o,j=inputs(['YF0001_上海语荻光电有限公司','上海语荻光电有限公司'])
    j.loc[1,['借方金额','贷方金额']]=[0,30]
    source=tmp_path/'原账.xlsx';j.to_excel(source,sheet_name='序时账',index=False)
    digest=hashlib.sha256(source.read_bytes()).digest()
    spec=importlib.util.spec_from_file_location('names_cli',Path(__file__).resolve().parents[1]/'scripts/计算往来账龄.py')
    cli=importlib.util.module_from_spec(spec);spec.loader.exec_module(cli)
    output=tmp_path/'账龄.xlsx'
    result=cli.execute({'input':str(source),'journal_sheet':'序时账','zero_opening':True,
        'cutoffs':['2025-12-31'],'include_receivable_bills':False,'output':str(output)})
    assert result['status']=='complete'
    book=load_workbook(output,data_only=True)
    def records(title):
        values=iter(book[title].values);headers=next(values)
        return [dict(zip(headers,row)) for row in values]
    final=[r for r in records('账龄计算结果') if r['科目名称']=='应收账款']
    assert len(final)==1 and final[0]['明细名称']=='上海语荻光电有限公司' and final[0]['期末余额']==70
    assert records('名称归一对照')[0]['原始称谓']=='YF0001_上海语荻光电有限公司'
    assert records('计算用序时账')[0]['原始往来单位']=='YF0001_上海语荻光电有限公司'
    book.close();assert hashlib.sha256(source.read_bytes()).digest()==digest


def test_已知不同登记主体即使给了合并决定也拒绝():
    a,b='成都示例有限公司','成都示例股份有限公司'
    o,j=inputs([a,b])
    with pytest.raises(ValueError,match='不同主体'):
        normalize_accounts(o,j,{'normalization_registry':[
            {'name':a,'registered_name':a,'credit_code':'A','source':'登记甲'},
            {'name':b,'registered_name':b,'credit_code':'B','source':'登记乙'}],
            'normalization_decisions':[{'names':[a,b],'canonical':a,'basis':'错误合并决定'}]})


def test_保留上游原名同时以本次输入名称记对照():
    a='YF0001_成都示例有限公司'
    o,j=inputs([a]);j['原始往来单位']='最初原账称谓'
    result=normalize_accounts(o,j,{})
    assert result['journal'].iloc[0]['原始往来单位']=='最初原账称谓'
    assert result['changes'][0]['原始称谓']==a


def test_同名但不同登记代码的整组别名都保持区分():
    a,aa,b,bb='YF0001_成都示例有限公司','成都示例有限公司','YF0002_成都示例有限公司','YF0003_成都示例有限公司'
    o,j=inputs([a,aa,b,bb])
    result=normalize_accounts(o,j,{'normalization_registry':[
        {'name':a,'registered_name':aa,'credit_code':'A','source':'登记甲'},
        {'name':b,'registered_name':aa,'credit_code':'B','source':'登记乙'}],
        'normalization_decisions':[
            {'names':[a,aa],'canonical':aa,'basis':'原账往来编码甲对应同一主体'},
            {'names':[b,bb],'canonical':aa,'basis':'原账往来编码乙对应另一主体'}]})
    assert result['status']=='ready'
    names=list(result['journal']['明细名称'])
    assert names[0]==names[1] and names[2]==names[3] and names[0]!=names[2]


def test_搬移技能后名称统一仍能承接原应收账龄至票据(tmp_path):
    root=Path(__file__).resolve().parents[1];portable=tmp_path/'独立技能'
    shutil.copytree(root,portable,ignore=shutil.ignore_patterns('.git','__pycache__','.pytest_cache'))
    source=tmp_path/'原账.xlsx';output=tmp_path/'账龄.xlsx'
    pd.DataFrame([
        [date(2023,1,1),'销售','1122','应收账款','YF0001_上海语荻光电有限公司',100,0,'',''],
        [date(2025,6,1),'收票','1122','应收账款','上海语荻光电有限公司',0,100,'',''],
        [date(2025,6,1),'收票','1121','应收票据','上海语荻光电有限公司',100,0,'商业承兑','票1'],
    ],columns=['日期','凭证号','科目编码','科目名称','明细名称','借方金额','贷方金额','票据类型','票据号']).to_excel(source,index=False)
    params=tmp_path/'助手参数.json'
    params.write_text(json.dumps({'input':str(source),'zero_opening':True,
        'cutoffs':['2025-12-31'],'output':str(output)},ensure_ascii=False),encoding='utf-8')
    response=subprocess.run([sys.executable,'-I','-X','utf8',str(portable/'scripts/计算往来账龄.py'),
        'calculate','--config',str(params)],cwd=tmp_path,capture_output=True,text=True,encoding='utf-8')
    assert response.returncode==0,response.stderr
    assert json.loads(response.stdout)['status']=='complete'
    book=load_workbook(output,data_only=True)
    rows=iter(book['账龄计算结果'].values);headers=next(rows)
    bill=next(dict(zip(headers,row)) for row in rows if dict(zip(headers,row))['科目名称']=='应收票据')
    assert bill['期末余额']==100 and bill['2—3年']==100 and bill['1年以内']==0
    assert '名称归一对照' in book.sheetnames and '票据连续计龄依据' in book.sheetnames
    book.close()
