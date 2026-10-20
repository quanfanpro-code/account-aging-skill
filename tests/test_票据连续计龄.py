from datetime import date

import pandas as pd
import pytest

from src.core.receivable_bills import ReceivableBillEngine, BillConfirmationRequired

BUCKETS=[(0,365),(365,730),(730,9999)]
CUTOFF=date(2025,12,31)

def journal(rows):
    return pd.DataFrame(rows,columns=['原账行号','日期','凭证号','科目编码','科目名称','明细名称','借方金额','贷方金额','票据类型','票据号'])

def opening():
    return pd.DataFrame(columns=['科目编码','科目名称','明细名称','借方余额','贷方余额'])

def run(rows,**kwargs):
    engine=ReceivableBillEngine(CUTOFF,BUCKETS)
    return engine,engine.run(opening(),journal(rows),**kwargs)

def initial():
    return [[2,date(2023,1,1),'销售1','1122','应收账款','单位甲',200,0,'',''],
            [3,date(2025,1,1),'销售2','1122','应收账款','单位甲',100,0,'','']]

def test_商业承兑按原应收两批账龄承接且分别列示科目():
    engine,results=run(initial()+[
        [4,date(2025,6,1),'收票1','1121','应收票据','单位甲',250,0,'商业承兑','票1'],
        [5,date(2025,6,1),'收票1','1122','应收账款','单位甲',0,250,'','']])
    by={r.subject_name:r for r in results}
    assert by['应收账款'].closing_balance==50 and by['应收账款'].bucket_totals==[50,0,0]
    assert by['应收票据'].closing_balance==250 and by['应收票据'].bucket_totals==[50,0,200]
    assert sum(x['承接金额'] for x in engine.transfer_records)==250
    assert all(e.bill_transfer_note for e in by['应收票据'].entries)


def test_收票部分兑付转回应收再收款仍沿用原债权日期():
    engine,results=run(initial()[:1]+[
        [3,date(2024,1,1),'收票','1122','应收账款','单位甲',0,200,'',''],
        [4,date(2024,1,1),'收票','1121','应收票据','单位甲',200,0,'商业承兑','票1'],
        [5,date(2025,1,1),'兑付','1121','应收票据','单位甲',0,40,'商业承兑','票1'],
        [6,date(2025,2,1),'退票','1121','应收票据','单位甲',0,160,'商业承兑','票1'],
        [7,date(2025,2,1),'退票','1122','应收账款','单位甲',160,0,'',''],
        [8,date(2025,3,1),'现金回款','1122','应收账款','单位甲',0,20,'','']])
    ar=next(r for r in results if r.subject_name=='应收账款')
    assert ar.closing_balance==140 and ar.bucket_totals==[0,0,140]
    assert ar.entries[0].source_date==date(2023,1,1)
    assert len(engine.transfer_records)==2


def test_有票据号的兑付只核销对应票据():
    rows=initial()+[
        [4,date(2025,2,1),'收票1','1121','应收票据','单位甲',200,0,'商业承兑','票1'],
        [5,date(2025,2,1),'收票1','1122','应收账款','单位甲',0,200,'',''],
        [6,date(2025,3,1),'收票2','1121','应收票据','单位甲',100,0,'商业承兑','票2'],
        [7,date(2025,3,1),'收票2','1122','应收账款','单位甲',0,100,'',''],
        [8,date(2025,4,1),'兑付2','1121','应收票据','单位甲',0,100,'商业承兑','票2']]
    _,results=run(rows)
    bill=next(r for r in results if r.subject_name=='应收票据')
    assert bill.closing_balance==200 and bill.bucket_totals==[0,0,200]


@pytest.mark.parametrize('kind',['银行承兑',''])
def test_银行或不明票据类型不能自动套商业承兑规则(kind):
    rows=initial()+[
        [4,date(2025,6,1),'收票','1121','应收票据','单位甲',200,0,kind,'票1'],
        [5,date(2025,6,1),'收票','1122','应收账款','单位甲',0,200,'','']]
    with pytest.raises(BillConfirmationRequired) as error:
        run(rows)
    assert error.value.candidates[0]['转出行']==5
    decision={**error.value.candidates[0],'连续计龄':False,'依据':'已核实该银行承兑结算使原应收债权终止确认'}
    _,results=run(rows,decisions=[decision])
    assert next(r for r in results if r.subject_name=='应收票据').bucket_totals==[200,0,0]


def test_明确确认适用的银行票据也可承接但须留下依据():
    rows=initial()+[
        [4,date(2025,6,1),'收票','1121','应收票据','单位甲',200,0,'银行承兑','票1'],
        [5,date(2025,6,1),'收票','1122','应收账款','单位甲',0,200,'','']]
    _,results=run(rows,decisions=[{'转出行':5,'转入行':4,'金额':200,'连续计龄':True,'依据':'经核实原债权未终止确认，本次按连续计龄政策分析'}])
    assert next(r for r in results if r.subject_name=='应收票据').bucket_totals==[0,0,200]


def test_不同单位相同金额不承接():
    engine,results=run(initial()+[
        [4,date(2025,6,1),'业务','1121','应收票据','单位乙',200,0,'商业承兑','票1'],
        [5,date(2025,6,1),'业务','1122','应收账款','单位甲',0,200,'','']])
    assert not engine.transfer_records
    assert next(r for r in results if r.subject_name=='应收票据').bucket_totals==[200,0,0]


def test_截止日后的收票不能提前承接():
    _,results=run(initial()+[
        [4,date(2026,1,1),'收票','1121','应收票据','单位甲',200,0,'商业承兑','票1'],
        [5,date(2026,1,1),'收票','1122','应收账款','单位甲',0,200,'','']])
    assert len(results)==1 and results[0].closing_balance==300


def test_原应收余额不足不凭空生成旧账龄():
    with pytest.raises(ValueError,match='余额不足'):
        run(initial()[:1]+[
            [4,date(2025,6,1),'收票','1121','应收票据','单位甲',250,0,'商业承兑','票1'],
            [5,date(2025,6,1),'收票','1122','应收账款','单位甲',0,250,'','']])


def test_期初分档经收票继续老化():
    balances=pd.DataFrame([['1122','应收账款','单位甲',100,0,0,100]],
        columns=['科目编码','科目名称','明细名称','借方余额','贷方余额','一年以内','一至二年'])
    engine=ReceivableBillEngine(CUTOFF,BUCKETS,opening_date=date(2024,12,31),opening_aging_columns={(0,365):'一年以内',(366,730):'一至二年'})
    results=engine.run(balances,journal([
        [4,date(2025,6,1),'收票','1121','应收票据','单位甲',100,0,'商业承兑','票1'],
        [5,date(2025,6,1),'收票','1122','应收账款','单位甲',0,100,'','']]))
    assert next(r for r in results if r.subject_name=='应收票据').bucket_totals==[0,0,100]


def test_完整技能工作流票据单独列示并输出连续账龄依据(tmp_path):
    from src.core.workflow import AgingWorkflow
    from src.core.report import ReportGenerator
    from openpyxl import load_workbook
    frame=journal(initial()+[
        [4,date(2025,6,1),'收票','1121','应收票据','单位甲',250,0,'商业承兑','票1'],
        [5,date(2025,6,1),'收票','1122','应收账款','单位甲',0,250,'','']])
    flow=AgingWorkflow(CUTOFF,BUCKETS)
    results=flow.run(opening(),frame,include_receivable_bills=True)
    assert {r.subject_name:r.closing_balance for r in results}=={'应收账款':50,'应收票据':250}
    path=tmp_path/'技能账龄.xlsx'
    ReportGenerator().generate(path,results,{'buckets':BUCKETS,'source_results':flow.source_results,'bill_transfer_records':flow.bill_transfer_records})
    wb=load_workbook(path)
    assert '票据连续计龄依据' in wb.sheetnames
    detail=list(wb['期末余额构成及账龄依据'].values)
    bill_rows=[row for row in detail[1:] if row[2]=='应收票据']
    assert sum(row[6] for row in bill_rows)==250
    assert all('票据连续计龄' in row[-1] for row in bill_rows)
    assert wb['账龄计算结果'].sheet_properties.tabColor.rgb=='000070C0'
    assert wb['票据连续计龄依据'].sheet_properties.tabColor is None
    wb.close()


def test_连续计龄模式不丢失反向应收期初的老分档():
    from src.core.workflow import AgingWorkflow
    balances=pd.DataFrame([['1122','应收账款','单位甲',0,100,0,-100]],
        columns=['科目编码','科目名称','明细名称','借方余额','贷方余额','一年以内','一至二年'])
    flow=AgingWorkflow(CUTOFF,BUCKETS,opening_date=date(2024,12,31),opening_aging_columns={(0,365):'一年以内',(366,730):'一至二年'})
    results=flow.run(balances,journal([]),include_receivable_bills=True)
    assert len(results)==1 and results[0].subject_name=='预收账款'
    assert results[0].bucket_totals==[0,0,100]


def test_红字冲回收票保持原始债权来源():
    _,results=run(initial()[:1]+[
        [3,date(2024,1,1),'收票','1121','应收票据','单位甲',200,0,'商业承兑','票1'],
        [4,date(2024,1,1),'收票','1122','应收账款','单位甲',0,200,'',''],
        [5,date(2025,1,1),'冲回','1121','应收票据','单位甲',-200,0,'商业承兑','票1'],
        [6,date(2025,1,1),'冲回','1122','应收账款','单位甲',0,-200,'','']])
    assert next(r for r in results if r.subject_name=='应收账款').bucket_totals==[0,0,200]
