from datetime import date
import hashlib
import importlib.util
import json
from pathlib import Path
import shutil
import subprocess
import sys

import pandas as pd
import pytest
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('aging_cli',ROOT/'scripts/计算往来账龄.py')
module=importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


def inputs(tmp_path,kind='商业承兑'):
    source=tmp_path/'账表.xlsx'
    journal=pd.DataFrame([
        [date(2023,1,1),'销售1','1122','应收账款','单位甲',100,0,'',''],
        [date(2023,1,1),'销售2','1122','应收账款','单位甲',100,0,'',''],
        [date(2025,6,1),'收票','1122','应收账款','单位甲',0,200,'',''],
        [date(2025,6,1),'收票','1121','应收票据','单位甲',200,0,kind,'票1'],
    ],columns=['日期','凭证号','科目编码','科目名称','明细名称','借方金额','贷方金额','票据类型','票据号'])
    journal.to_excel(source,index=False,sheet_name='序时账')
    config={'input':str(source),'journal_sheet':'序时账','zero_opening':True,
        'cutoffs':['2024-12-31','2025-12-31'],'output':str(tmp_path/'最终账龄.xlsx')}
    return config


def data(sheet):
    rows=list(sheet.values)
    return [dict(zip(rows[0][1:],row[1:])) for row in rows[1:]]


def test_真实多时点导出不误删同日同金额的两个余额层(tmp_path):
    config=inputs(tmp_path)
    original=hashlib.sha256(Path(config['input']).read_bytes()).digest()
    assert module.execute(config,inspect_only=True)['status']=='ready'
    assert not Path(config['output']).exists()
    result=module.execute(config)
    assert result['status']=='complete'
    book=load_workbook(config['output'],data_only=True)
    details=data(book['期末余额构成及账龄依据'])
    assert len(details)==4
    for cutoff,category in [('2024-12-31','应收账款'),('2025-12-31','应收票据')]:
        sheet=book[cutoff+'账龄结果']
        assert sheet.sheet_properties.tabColor.rgb=='000070C0'
        final=next(row for row in data(sheet) if row['科目名称']==category)
        layers=[row for row in details if str(row['截止日期'])[:10]==cutoff and row['科目名称']==category]
        assert len(layers)==2 and sum(row['金额'] for row in layers)==final['期末余额']==200
    assert all(s.sheet_properties.tabColor is None for s in book if not s.title.endswith('账龄结果'))
    book.close()
    assert hashlib.sha256(Path(config['input']).read_bytes()).digest()==original


def test_不明票据对话确认后才导出且保留依据(tmp_path):
    config=inputs(tmp_path,kind='银行承兑')
    response=module.execute(config,inspect_only=True)
    assert response['status']=='needs_bill_decisions'
    assert not Path(config['output']).exists()
    config['bill_decisions']=[{**response['票据确认'][0],'连续计龄':False,'依据':'原债权已终止确认，按已核实的会计政策自收票日计龄'}]
    assert module.execute(config)['status']=='complete'
    book=load_workbook(config['output'],data_only=True)
    row=next(row for row in data(book['2025-12-31账龄结果']) if row['科目名称']=='应收票据')
    assert row['1年以内']==200
    assert '已终止确认' in data(book['特殊事项'])[0]['处理依据']
    book.close()


def test_无期初不能未经确认假定为零(tmp_path):
    config=inputs(tmp_path)
    config.pop('zero_opening')
    with pytest.raises(ValueError,match='确认历史完整'):
        module.execute(config,inspect_only=True)


def test_自包含目录搬移后仍可调用不依赖原项目(tmp_path):
    portable=tmp_path/'便携技能'
    shutil.copytree(ROOT,portable,ignore=shutil.ignore_patterns('.git','__pycache__','.pytest_cache'))
    config=inputs(tmp_path)
    params=tmp_path/'agent参数.json'
    params.write_text(json.dumps(config,ensure_ascii=False),encoding='utf-8')
    result=subprocess.run([sys.executable,'-X','utf8',str(portable/'scripts/计算往来账龄.py'),'calculate','--config',str(params)],
        cwd=tmp_path,capture_output=True,text=True,encoding='utf-8')
    assert result.returncode==0,result.stderr
    assert json.loads(result.stdout)['status']=='complete'
    assert Path(config['output']).is_file()
