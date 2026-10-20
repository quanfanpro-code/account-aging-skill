# -*- coding: utf-8 -*-
import os
from datetime import date
import openpyxl
import pytest
from src.core.fifo_engine import FIFOEngine
from src.core.report import ReportGenerator


def test_生成Excel_只有一套最终账龄及追溯记录(tmp_path):
    result=FIFOEngine(date(2025,6,1),[(0,365)]).calculate('1122','应收账款','客户A',100,[])
    target=tmp_path/'最终结果.xlsx'
    ReportGenerator().generate(target,[result],{'buckets':[(0,365)]})
    wb=openpyxl.load_workbook(target)
    assert wb.sheetnames==['账龄计算结果','期末余额构成及账龄依据','原账勾稽','处理确认记录','内部转科目记录','重分类记录','对抵记录','参数说明']
    assert wb['账龄计算结果']['E2'].value==100
    assert wb['账龄计算结果']['G2'].value=='资产'
    assert wb['账龄计算结果'].sheet_properties.tabColor.rgb=='000070C0'
    assert all(sheet.sheet_properties.tabColor is None for sheet in wb if sheet.title!='账龄计算结果')
    wb.close()


def test_最终结果资产负债分别合计(tmp_path):
    from src.core.models import AgingResult
    results=[AgingResult('1122','应收账款','A','debit',100,0,0,100,bucket_totals=[100]),AgingResult('2202','应付账款','B','credit',50,0,0,50,bucket_totals=[50])]
    target=tmp_path/'合计.xlsx'
    ReportGenerator().generate(target,results,{'buckets':[(0,365)]})
    wb=openpyxl.load_workbook(target,data_only=True);s=wb['账龄计算结果']
    assert s['B4'].value=='资产合计' and s['F4'].value==100
    assert s['B5'].value=='负债合计' and s['F5'].value==50
    wb.close()


def test_计算用期初保留原分档与名称(tmp_path):
    import pandas as pd
    source=pd.DataFrame({'明细名称':['原名称'],'期初一年内':[100]})
    target=tmp_path/'期初来源.xlsx'
    ReportGenerator().generate(target,[],{'buckets':[(0,365)],'opening_source':source})
    wb=openpyxl.load_workbook(target,data_only=True)
    assert wb['计算用期初']['B2'].value=='原名称' and wb['计算用期初']['C2'].value==100
    wb.close()


def test_自定义分档标签_出现在Excel表头(tmp_path):
    """传入 bucket_labels 时，Excel 表头应使用自定义标签而非 '0-90天'"""
    engine = FIFOEngine(
        cutoff_date=date(2025, 6, 1),
        buckets=[(0, 90), (90, 365), (365, 9999)],
        direction="debit",
    )
    result = engine.calculate(
        subject_code="1122", subject_name="应收账款", detail_name="客户A",
        opening_balance=1000.0,
        transactions=[{"date": date(2024, 6, 1), "debit": 500.0, "credit": 0.0}],
    )
    output = tmp_path / "label_test.xlsx"
    gen = ReportGenerator()
    gen.generate(
        output_path=output, aging_results=[result],
        config={"cutoff_date": date(2025, 6, 1),
                "buckets": [(0, 90), (90, 365), (365, 9999)],
                "direction": "debit",
                "bucket_labels": ["3个月内", "4-12月", "1年以上"]},
    )
    wb = openpyxl.load_workbook(output)
    sheet = wb["账龄计算结果"]
    # 表头应包含自定义标签
    header_row = [str(cell.value) for cell in sheet[1]]
    header_text = " ".join(header_row)
    assert "3个月内" in header_text
    assert "4-12月" in header_text
    assert "1年以上" in header_text
    # 不应包含原始的天数格式
    assert "0-90天" not in header_text


def test_无自定义标签时回退到天数格式(tmp_path):
    """不传 bucket_labels 时，表头仍用 '0-90天' 格式"""
    engine = FIFOEngine(
        cutoff_date=date(2025, 6, 1),
        buckets=[(0, 90), (90, 365)],
        direction="debit",
    )
    result = engine.calculate(
        subject_code="1122", subject_name="应收", detail_name="A",
        opening_balance=0.0,
        transactions=[{"date": date(2024, 6, 1), "debit": 1000.0, "credit": 0.0}],
    )
    output = tmp_path / "nolabel.xlsx"
    gen = ReportGenerator()
    gen.generate(
        output_path=output, aging_results=[result],
        config={"cutoff_date": date(2025, 6, 1), "buckets": [(0, 90), (90, 365)], "direction": "debit"},
    )
    wb = openpyxl.load_workbook(output)
    sheet = wb["账龄计算结果"]
    header_text = " ".join(str(cell.value) for cell in sheet[1])
    assert "0-90天" in header_text
    assert "90-365天" in header_text


def test_参数说明_包含期初日期与模式说明(tmp_path):
    engine = FIFOEngine(
        cutoff_date=date(2025, 6, 1),
        buckets=[(0, 90), (90, 365)],
        direction="debit",
    )
    result = engine.calculate(
        subject_code="1122", subject_name="应收", detail_name="A",
        opening_balance=1000.0,
        transactions=[],
        opening_date=date(2024, 1, 1),
    )
    output = tmp_path / "config.xlsx"
    ReportGenerator().generate(
        output_path=output,
        aging_results=[result],
        config={
            "cutoff_date": date(2025, 6, 1),
            "buckets": [(0, 90), (90, 365)],
            "direction": "debit",
            "opening_date": date(2024, 1, 1),
            "aging_mode": "铁板模式",
            "bucket_boundary_rule": "边界点归前档",
        },
    )
    wb = openpyxl.load_workbook(output)
    sheet = wb["参数说明"]
    values = " ".join(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value)
    assert "期初余额日期" in values
    assert "铁板模式" in values
    assert "边界点归前档" in values


def test_未经重分类的反向余额不得直接作为最终账龄导出(tmp_path):
    result=FIFOEngine(date(2025,6,1),[(0,365)]).calculate('1122','应收账款','A',100,[{'date':date(2025,1,1),'debit':0,'credit':150}])
    with pytest.raises(ValueError,match='重分类'):
        ReportGenerator().generate(tmp_path/'无效结果.xlsx',[result],{'buckets':[(0,365)]})


def test_原账勾稽采用重分类对抵前来源(tmp_path):
    from dataclasses import replace
    original=FIFOEngine(date(2025,6,1),[(0,365)]).calculate('1122','应收账款','A',100,[])
    final=replace(original,closing_balance=70,bucket_totals=[70])
    target=tmp_path/'原账核对.xlsx'
    ReportGenerator().generate(target,[final],{'buckets':[(0,365)],'source_results':[original]})
    wb=openpyxl.load_workbook(target,data_only=True)
    assert wb['原账勾稽']['H2'].value==100
    assert wb['账龄计算结果']['F2'].value==70
    assert wb['原账勾稽']['J2'].value==0
    wb.close()


def test_最终金额与分档勾稽不平衡拒绝输出(tmp_path):
    from src.core.models import AgingResult
    bad=AgingResult('1122','应收账款','A','debit',100,0,0,100,bucket_totals=[90])
    with pytest.raises(ValueError,match='不一致'):
        ReportGenerator().generate(tmp_path/'不平衡.xlsx',[bad],{'buckets':[(0,365)]})


def test_显式空白或重复账龄名称拒绝导出():
    with pytest.raises(ValueError, match="不能为空"):
        ReportGenerator._get_bucket_labels({
            "buckets": [(0, 90), (90, 180)], "bucket_labels": ["", ""],
        })
    with pytest.raises(ValueError, match="不能重复"):
        ReportGenerator._get_bucket_labels({
            "buckets": [(0, 90), (90, 180)], "bucket_labels": ["半年内", "半年内"],
        })


def test_未提供名称时按档位生成非空默认名称():
    assert ReportGenerator._get_bucket_labels({
        "buckets": [(0, 90), (730, 9999)],
    }) == ["0-90天", "730天以上"]


def test_无反向余额也保留反向余额列():
    engine = FIFOEngine(
        cutoff_date=date(2025, 6, 1), buckets=[(0, 365)], direction="debit",
    )
    result = engine.calculate(
        subject_code="1122", subject_name="应收账款", detail_name="客户A",
        opening_balance=100.0, transactions=[],
    )
    summary = ReportGenerator()._build_summary(
        [result], {"buckets": [(0, 365)], "bucket_labels": ["1年以内"]},
    )
    assert list(summary.columns) == [
        "科目编码", "科目名称", "明细名称", "1年以内",
        "反向余额（待核对）", "合计",
    ]
    assert summary.iloc[0]["反向余额（待核对）"] == 0.0


def test_参数说明使用固定方向和反向余额新口径():
    config_table = ReportGenerator()._build_config({
        "buckets": [(0, 365)],
        "bucket_labels": ["1年以内"],
        "direction_policy": "八类往来按科目类别固定方向",
        "negative_balance_policy": "反向余额以负数单列展示，不改变科目所属汇总表",
    })
    text = " ".join(str(value) for value in config_table.to_numpy().ravel())
    assert "八类往来按科目类别固定方向" in text
    assert "反向余额以负数单列展示" in text
    assert "默认方向" not in text
